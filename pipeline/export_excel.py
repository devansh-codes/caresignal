"""Build web/public/downloads/CareSignal_Workbook.xlsx — a live, formula-driven workbook.

Every derived number on the calculator sheets is an Excel formula over the Data sheet, so a reader can
audit the HAC methodology and run what-if scenarios without Python. Python-computed values are written
alongside (grey "Python" columns) purely as a cross-check.

Compatibility: lookups use XLOOKUP (Excel 2021/365, LibreOffice 24.8+) wrapped in IFERROR with an
INDEX/MATCH fallback, so they still resolve in older Excel. LET is used only where an equivalent plain
formula is supplied as the IFERROR fallback. Excel-2010+ functions are written with their `_xlfn.`
storage prefix (PERCENTILE.INC, STDEV.P, PERCENTRANK.INC) as the OOXML spec requires.

Run: python -m pipeline.export_excel
"""

from __future__ import annotations

import json
import math
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
PROCESSED = ROOT / "data" / "processed"
RAW = ROOT / "data" / "raw"
OUT = ROOT / "web" / "public" / "downloads" / "CareSignal_Workbook.xlsx"

HAC_MEASURES = ["psi90", "clabsi", "cauti", "ssi", "cdi", "mrsa"]
HAI_MEASURES = ["clabsi", "cauti", "ssi", "cdi", "mrsa"]
MEASURE_LABEL = {
    "psi90": "PSI 90 Patient Safety composite",
    "clabsi": "CLABSI (central line infection)",
    "cauti": "CAUTI (catheter urinary infection)",
    "ssi": "SSI (colon + abdominal hysterectomy)",
    "cdi": "C. difficile (CDI)",
    "mrsa": "MRSA bacteremia",
}
MEASURE_CATEGORY = {
    "psi90": "Shared",
    "clabsi": "Nursing-driven",
    "cauti": "Nursing-driven",
    "ssi": "Other",
    "cdi": "Shared",
    "mrsa": "Shared",
}

# ---- Styling ---------------------------------------------------------------------------------------
TEAL = "0B6E8A"
HEADER_FILL = PatternFill("solid", fgColor=TEAL)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color=TEAL)
SUB_FONT = Font(italic=True, size=10, color="546376")
SECTION_FONT = Font(bold=True, size=12, color=TEAL)
INPUT_FILL = PatternFill("solid", fgColor="DDEBF7")
INPUT_FONT = Font(color="0000FF", bold=True)  # finance convention: blue = user input
HELPER_FONT = Font(color="808080", italic=True)
PY_FONT = Font(color="808080")
CALC_FILL = PatternFill("solid", fgColor="F2F5F8")
RED_FILL = PatternFill("solid", fgColor="F8D7DA")
RED_FONT = Font(color="9C0006", bold=True)
GREEN_FILL = PatternFill("solid", fgColor="D9EFE1")
GREEN_FONT = Font(color="1E7A46", bold=True)
AMBER_FILL = PatternFill("solid", fgColor="FFF1D6")
THIN = Side(style="thin", color="DEE4EC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

F_DEC4 = "0.0000"
F_DEC3 = "0.000"
F_DEC2 = "0.00"
F_INT = "#,##0"
F_USD = '"$"#,##0'
F_PCT_PTS = "0.0"
F_PCT = "0.0%"


def _header(ws, row: int, col: int, values: list[str]) -> None:
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _title(ws, text: str, sub: str | None = None) -> None:
    ws["B2"] = text
    ws["B2"].font = TITLE_FONT
    if sub:
        ws["B3"] = sub
        ws["B3"].font = SUB_FONT
    ws.column_dimensions["A"].width = 2
    ws.sheet_view.showGridLines = False


def _input(cell) -> None:
    cell.fill, cell.font, cell.border = INPUT_FILL, INPUT_FONT, BOX


def _nan(v):
    if v is None:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    if v is pd.NA or v is pd.NaT:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(v, "item"):  # numpy scalar -> python
        v = v.item()
    return v


# ---- Inputs ----------------------------------------------------------------------------------------
@dataclass
class Inputs:
    data: pd.DataFrame  # merged hospital metrics + HAC values
    params: dict
    dq: pd.DataFrame | None
    measure_map: list[dict]
    sources: list[dict]
    built_at: str


def _yes_no(v):
    v = _nan(v)
    if v is None:
        return None
    if isinstance(v, str):
        return v if v in ("Yes", "No") else ("Yes" if v.lower() in ("true", "yes", "y", "1") else "No")
    return "Yes" if bool(v) else "No"


def load_inputs() -> Inputs:
    hm = pd.read_parquet(PROCESSED / "hospital_metrics.parquet")
    hac = pd.read_parquet(PROCESSED / "hac_rebuild.parquet")
    params = json.loads((PROCESSED / "hac_params.json").read_text())
    hm["ccn"] = hm["ccn"].astype(str).str.zfill(6)
    hac["ccn"] = hac["ccn"].astype(str).str.zfill(6)

    # HAC inputs: the value CMS scored. A value only enters the winsorization population when the
    # rebuild produced a z for it (footnoted/suppressed values are excluded the same way).
    keep = ["ccn", "is_maryland", "total_official", "total_rebuilt"]
    for m in HAC_MEASURES:
        keep += [f"{m}_value", f"{m}_z_official", f"{m}_z_rebuilt"]
    keep = [c for c in keep if c in hac.columns]
    h = hac[keep].copy()
    for m in HAC_MEASURES:
        zr, zo = f"{m}_z_rebuilt", f"{m}_z_official"
        mask = h[zr].notna() if zr in h else h[zo].notna()
        h[f"{m}_input"] = h[f"{m}_value"].where(mask)
        # No value but a published z (footnotes 11/18, data not submitted): CMS assigns the maximum z,
        # i.e. the value is treated as the 95th-percentile cap. These are excluded from the distribution.
        h[f"{m}_maxz"] = (h[f"{m}_value"].isna() & mask).map({True: "Yes", False: None})
    h = h.rename(columns={f"{m}_value": f"{m}_value_hac" for m in HAC_MEASURES})

    drop_dupe = [c for c in hm.columns if c in h.columns and c != "ccn"]
    df = hm.merge(h, on="ccn", how="outer", suffixes=("", "_hac"))
    for c in drop_dupe:
        if f"{c}_hac" in df:
            df[c] = df[c].combine_first(df[f"{c}_hac"])
            df = df.drop(columns=[f"{c}_hac"])
    for c in ("name", "city", "state", "hospital_type", "peer_key", "peer_label"):
        if c not in df:
            df[c] = None
    df["name"] = df["name"].fillna("(not in General Information file)")
    df = df.sort_values("ccn").reset_index(drop=True)

    dq_path = PROCESSED / "dq_log.csv"
    dq = pd.read_csv(dq_path, dtype=str) if dq_path.exists() else None

    mm_path = ROOT / "config" / "measure_map.yaml"
    measure_map: list[dict] = []
    if mm_path.exists():
        raw = yaml.safe_load(mm_path.read_text()) or {}
        items = raw.get("measures", raw) if isinstance(raw, dict) else raw
        if isinstance(items, dict):
            for k, v in items.items():
                measure_map.append({"id": k, **(v if isinstance(v, dict) else {"category": v})})
        elif isinstance(items, list):
            measure_map = [i for i in items if isinstance(i, dict)]

    sources = []
    for mp in (RAW / "manifest.json", ROOT / "web" / "public" / "data" / "manifest.json"):
        if mp.exists():
            m = json.loads(mp.read_text())
            src = m.get("sources", m) if isinstance(m, dict) else m
            if isinstance(src, dict):
                src = [{"id": k, **v} for k, v in src.items()]
            sources = src
            break

    built_at = datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")
    return Inputs(df, params, dq, measure_map, sources, built_at)


# ---- Data sheet column spec ------------------------------------------------------------------------
@dataclass
class Col:
    key: str
    header: str
    fmt: str | None = None
    width: float = 12
    formula: bool = False


def data_columns(df: pd.DataFrame) -> list[Col]:
    cols = [
        Col("ccn", "CCN", "@", 9),
        Col("label", "Hospital (CCN — Name)", None, 44),
        Col("name", "Hospital", None, 34),
        Col("city", "City", None, 16),
        Col("state", "State", None, 7),
        Col("hospital_type", "Hospital type", None, 22),
        Col("peer_key", "Peer group key", None, 28),
        Col("peer_label", "Peer group", None, 34),
        Col("bed_band", "Bed band", None, 10),
        Col("teaching", "Teaching", None, 10),
        Col("urban_rural", "Urban/rural", None, 10),
        Col("region", "Region", None, 11),
        Col("beds", "Beds", F_INT, 8),
        Col("is_maryland", "Maryland (waived)", None, 10),
        Col("hac_total", "Total HAC (official)", F_DEC4, 11),
        Col("payment_reduction", "Payment reduction (official)", None, 11),
        Col("hac_gap", "Gap to cutoff (Python)", F_DEC4, 11),
        Col("hac_status", "HAC status", None, 11),
        Col("hac_percentile", "HAC national percentile", F_PCT_PTS, 10),
    ]
    for m in HAC_MEASURES:
        cols.append(Col(f"{m}_input", f"{m.upper()} value (scored)", F_DEC4 if m == "psi90" else F_DEC3, 11))
    for m in HAC_MEASURES:
        cols.append(Col(f"{m}_z_official", f"{m.upper()} z (official)", F_DEC4, 10))
    for m in HAC_MEASURES:
        cols.append(Col(f"{m}_maxz", f"{m.upper()} max z assigned (not submitted)", None, 10))
    for m in HAI_MEASURES:
        cols.append(Col(f"{m}_predicted", f"{m.upper()} predicted infections", F_DEC3, 11))
    for m in HAC_MEASURES:
        cols.append(Col(f"ptp_{m}_events", f"{m.upper()} events to prevent (Python)", F_INT, 11))
    cols += [
        Col("tipping_measure", "Tipping measure", None, 10),
        Col("tipping_events", "Tipping events", F_INT, 9),
        Col("min_cushion_events", "Min cushion events", F_INT, 9),
        Col("traj_total", "Trajectory HAC (proxy)", F_DEC4, 11),
        Col("traj_status", "Trajectory status", None, 11),
        Col("nurse_comm_top_box", "Nurse comm. 'Always' %", F_PCT_PTS, 10),
        Col("nurse_comm_pctl_nat", "Nurse comm. national pctl (Python)", F_PCT_PTS, 10),
        Col("nurse_comm_pctl_peer", "Nurse comm. peer pctl (Python)", F_PCT_PTS, 10),
        Col("discharge_info_yes", "Discharge info 'Yes' %", F_PCT_PTS, 10),
        Col("care_transition", "Care transition %", F_PCT_PTS, 10),
        Col("hcahps_star", "HCAHPS summary star", "0", 8),
        Col("err_ami", "ERR AMI", F_DEC4, 9),
        Col("err_hf", "ERR HF", F_DEC4, 9),
        Col("err_pn", "ERR PN", F_DEC4, 9),
        Col("err_copd", "ERR COPD", F_DEC4, 9),
        Col("err_cabg", "ERR CABG", F_DEC4, 9),
        Col("err_hip_knee", "ERR Hip/Knee", F_DEC4, 9),
        Col("n_err_above_1", "# ERR > 1", F_INT, 8),
        Col("exposure_usd", "Est. 1% HAC exposure ($)", F_USD, 13),
        Col("vbp_tps", "VBP Total Performance Score", F_DEC2, 10),
    ]
    # Live formula block: winsorized values, rebuilt z-scores (contiguous), total, total excl. Maryland.
    for m in HAC_MEASURES:
        cols.append(Col(f"{m}_w", f"{m.upper()} winsorized (formula)", F_DEC4, 11, True))
    for m in HAC_MEASURES:
        cols.append(Col(f"{m}_zcalc", f"{m.upper()} z (formula)", F_DEC4, 10, True))
    cols.append(Col("total_calc", "Total HAC (formula)", F_DEC4, 11, True))
    cols.append(Col("total_calc_nonmd", "Total HAC excl. MD (formula)", F_DEC4, 11, True))
    return [c for c in cols if c.formula or c.key in df.columns or c.key == "label"]


class Book:
    """Holds sheet handles and column-letter lookups so formulas can be generated by key."""

    def __init__(self, inp: Inputs):
        self.inp = inp
        self.df = inp.data
        self.wb = Workbook()
        self.n = len(self.df)
        self.last = self.n + 1  # last data row on Data sheet
        self.cols = data_columns(self.df)
        self.letter = {c.key: get_column_letter(i + 1) for i, c in enumerate(self.cols)}
        # Fixed cell addresses shared across sheets (set as sheets are laid out).
        self.sel_ccn = "'Hospital Lookup'!$C$6"
        self.sel_row = "'Hospital Lookup'!$C$7"
        self.calc_param_row = {m: 7 + i for i, m in enumerate(HAC_MEASURES)}  # HAC Calculator rows
        self.cutoff_cell = "'HAC Calculator'!$D$15"

    def rng(self, key: str) -> str:
        L = self.letter[key]
        return f"Data!${L}$2:${L}${self.last}"

    def has(self, key: str) -> bool:
        return key in self.letter

    def param(self, m: str, what: str) -> str:
        col = {"lo": "M", "hi": "N", "mean": "O", "sd": "P", "n": "C"}[what]  # "Active" columns
        return f"'HAC Calculator'!${col}${self.calc_param_row[m]}"

    def lookup(self, key: str) -> str:
        """XLOOKUP with INDEX/MATCH fallback; blank -> em dash so missing never reads as 0."""
        ccn, rng = self.rng("ccn"), self.rng(key)
        x = f"_xlfn.XLOOKUP({self.sel_ccn},{ccn},{rng})"
        im = f"INDEX({rng},{self.sel_row})"
        return f'=IFERROR(IF(LEN({x})=0,"—",{x}),IFERROR(IF(LEN({im})=0,"—",{im}),"—"))'

    def lookup_num(self, key: str) -> str:
        """Numeric lookup returning "" (not —) when missing, for use inside downstream math."""
        rng = self.rng(key)
        im = f"INDEX({rng},{self.sel_row})"
        return f'IFERROR(IF(ISNUMBER({im}),{im},""),"")'


# ---- Sheets ----------------------------------------------------------------------------------------
def build_data(b: Book) -> None:
    ws = b.wb.create_sheet("Data")
    for i, c in enumerate(b.cols, start=1):
        ws.cell(row=1, column=i, value=c.header)
        ws.column_dimensions[get_column_letter(i)].width = c.width
    _header(ws, 1, 1, [c.header for c in b.cols])
    ws.row_dimensions[1].height = 45

    L = b.letter
    zc_first, zc_last = L[f"{HAC_MEASURES[0]}_zcalc"], L[f"{HAC_MEASURES[-1]}_zcalc"]
    df = b.df
    for r, rec in enumerate(df.to_dict("records"), start=2):
        for i, c in enumerate(b.cols, start=1):
            cell = ws.cell(row=r, column=i)
            if c.formula:
                if c.key.endswith("_w"):
                    m = c.key[:-2]
                    v = f"{L[f'{m}_input']}{r}"
                    clip = f'MIN(MAX({v},{b.param(m, "lo")}),{b.param(m, "hi")})'
                    if f"{m}_maxz" in L:
                        flag = f"{L[f'{m}_maxz']}{r}"
                        cell.value = f'=IF(ISNUMBER({v}),{clip},IF({flag}="Yes",{b.param(m, "hi")},""))'
                    else:
                        cell.value = f'=IF(ISNUMBER({v}),{clip},"")'
                elif c.key.endswith("_zcalc"):
                    m = c.key[:-6]
                    w = f"{L[f'{m}_w']}{r}"
                    cell.value = f'=IF(ISNUMBER({w}),({w}-{b.param(m, "mean")})/{b.param(m, "sd")},"")'
                elif c.key == "total_calc":
                    # Equal weights across the measures the hospital has = mean of available z-scores.
                    z = f"{zc_first}{r}:{zc_last}{r}"
                    cell.value = f'=IF(COUNT({z})=0,"",AVERAGE({z}))'
                elif c.key == "total_calc_nonmd":
                    t = f"{L['total_calc']}{r}"
                    md = f"{L['is_maryland']}{r}" if "is_maryland" in L else '"No"'
                    cell.value = f'=IF(OR({md}="Yes",NOT(ISNUMBER({t}))),"",{t})'
                cell.fill = CALC_FILL
            elif c.key == "label":
                cell.value = f"{rec['ccn']} — {rec.get('name') or ''}"
            elif c.key in ("is_maryland", "payment_reduction"):
                cell.value = _yes_no(rec.get(c.key))
            else:
                cell.value = _nan(rec.get(c.key))
            if c.fmt:
                cell.number_format = c.fmt
    last_col = get_column_letter(len(b.cols))
    t = Table(displayName="tblHospitals", ref=f"A1:{last_col}{b.last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = "C2"
    # Highlight hospitals above the cutoff (red reserved for "above cutoff").
    for key in ("hac_total", "total_calc"):
        if key in L:
            col = L[key]
            ws.conditional_formatting.add(
                f"{col}2:{col}{b.last}",
                FormulaRule(formula=[f"AND(ISNUMBER({col}2),{col}2>{b.cutoff_cell})"], fill=RED_FILL, font=RED_FONT),
            )
    b.wb.defined_names["DataCCN"] = DefinedName("DataCCN", attr_text=b.rng("ccn"))
    b.wb.defined_names["DataLabels"] = DefinedName("DataLabels", attr_text=b.rng("label"))


WHATIF_FIRST = 27  # first measure row of the what-if table on 'HAC Calculator'
SRC_CAL, SRC_FORMULA = "CMS-calibrated", "Formula rebuild"


def _pct_formula(values: str, p: float, method: str, cond: str) -> tuple[str, bool]:
    """Percentile over `values` restricted to `cond`, reproducing numpy's `method`. Returns (formula, exact)."""
    arr = f'IF({cond},{values},"")'
    k = f"{p}*(SUMPRODUCT(--({cond}))-1)"
    if method == "linear":
        return f"_xlfn.PERCENTILE.INC({arr},{p})", True
    if method == "nearest":  # numpy rounds half to even; ROUND differs only at an exact .5 rank
        return f"SMALL({arr},ROUND({k},0)+1)", True
    if method == "lower":
        return f"SMALL({arr},INT({k})+1)", True
    if method == "higher":
        return f"SMALL({arr},-INT(-{k})+1)", True
    if method == "midpoint":
        return f"(SMALL({arr},INT({k})+1)+SMALL({arr},-INT(-{k})+1))/2", True
    return f"_xlfn.PERCENTILE.INC({arr},{p})", False


def build_calculator(b: Book) -> None:
    ws = b.wb.create_sheet("HAC Calculator")
    _title(
        ws,
        "HAC Calculator — FY2026 CMS methodology, rebuilt with live formulas",
        "Blue cells are inputs. Grey columns are pipeline constants, shown for audit.",
    )
    p = b.inp.params
    conv = p.get("conventions", {})
    method = conv.get("percentile_method", "linear")
    ddof = int(conv.get("sd_ddof", 0))
    include_md = bool(conv.get("include_maryland", True))
    sd_fn = "_xlfn.STDEV.S" if ddof == 1 else "_xlfn.STDEV.P"
    cal = p.get("measures", {})
    indep = p.get("independent", cal)

    ws["B4"] = (
        f"Steps 1–3: winsorize each measure at its 5th/95th percentile across hospitals with a result "
        f"({'incl.' if include_md else 'excl.'} Maryland), then z = (winsorized − mean) ÷ SD. The formula columns use "
        f"the convention that best reproduces CMS's published z-scores: percentile method '{method}', "
        f"{'sample SD (STDEV.S)' if ddof == 1 else 'population SD (STDEV.P)'}. 'CMS-calibrated' parameters are recovered "
        f"exactly from the official W Z scores and match the CareSignal site."
    )
    ws["B4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("B4:P4")
    ws.row_dimensions[4].height = 46
    ws["B5"] = "Parameter source ▸"
    ws["B5"].font = Font(bold=True)
    ws["C5"] = SRC_CAL
    _input(ws["C5"])
    ws.merge_cells("C5:D5")
    dv = DataValidation(type="list", formula1=f'"{SRC_CAL},{SRC_FORMULA}"', allow_blank=False)
    dv.add("C5")
    ws.add_data_validation(dv)
    ws["E5"] = "Drives the Data sheet z-scores, Total HAC, cutoff, what-if and Patients to Prevent."
    ws["E5"].font = SUB_FONT

    _header(
        ws,
        6,
        2,
        [
            "Measure",
            "N in distribution",
            "P5 (formula)",
            "P95 (formula)",
            "Mean of winsorized (formula)",
            f"SD of winsorized ({'STDEV.S' if ddof == 1 else 'STDEV.P'}, formula)",
            "Max |Δ| formula vs pipeline",
            "Calibrated P5",
            "Calibrated P95",
            "Calibrated mean",
            "Calibrated SD",
            "Active P5",
            "Active P95",
            "Active mean",
            "Active SD",
        ],
    )
    md_rng = b.rng("is_maryland") if b.has("is_maryland") else None
    exact_all = True
    for m in HAC_MEASURES:
        r = b.calc_param_row[m]
        R = b.rng(f"{m}_input")
        cond = f"ISNUMBER({R})" + ("" if include_md or not md_rng else f'*({md_rng}<>"Yes")')
        ws.cell(row=r, column=2, value=MEASURE_LABEL[m])
        ws.cell(row=r, column=3, value=f"=SUMPRODUCT(--({cond}))").number_format = F_INT
        lo_f, ex1 = _pct_formula(R, 0.05, method, cond)
        hi_f, ex2 = _pct_formula(R, 0.95, method, cond)
        exact_all &= ex1 and ex2
        ws[f"D{r}"] = ArrayFormula(f"D{r}", "=" + lo_f)
        ws[f"E{r}"] = ArrayFormula(f"E{r}", "=" + hi_f)
        clip = f"IF({R}>E{r},E{r},IF({R}<D{r},D{r},{R}))"
        ws[f"F{r}"] = ArrayFormula(f"F{r}", f'=AVERAGE(IF({cond},{clip},""))')
        ws[f"G{r}"] = ArrayFormula(f"G{r}", f'={sd_fn}(IF({cond},{clip},""))')
        pi = indep.get(m, {})
        diffs = [
            f"ABS({c}{r}-{_nan(pi.get(k)) if _nan(pi.get(k)) is not None else 0})"
            for c, k in zip("DEFG", ("winsor_low", "winsor_high", "mean", "sd"), strict=False)
        ]
        ws[f"H{r}"] = f"=MAX({','.join(diffs)})"
        pc = cal.get(m, {})
        for c, k in zip("IJKL", ("winsor_low", "winsor_high", "mean", "sd"), strict=False):
            ws[f"{c}{r}"] = _nan(pc.get(k))
            ws[f"{c}{r}"].font = PY_FONT
        for act, fcol, ccol in zip("MNOP", "DEFG", "IJKL", strict=False):
            ws[f"{act}{r}"] = f'=IF($C$5="{SRC_FORMULA}",{fcol}{r},{ccol}{r})'
            ws[f"{act}{r}"].font = Font(bold=True)
        for col in "DEFGHIJKLMNOP":
            ws[f"{col}{r}"].number_format = "0.000000"
        for col in "BCDEFGHIJKLMNOP":
            ws[f"{col}{r}"].border = BOX
    if not exact_all:
        ws["B13"] = f"Note: percentile method '{method}' has no exact Excel equivalent; PERCENTILE.INC is used."
        ws["B13"].font = SUB_FONT

    tot_nonmd = b.rng("total_calc_nonmd")
    cut_f, _ = _pct_formula(tot_nonmd, 0.75, method, f"ISNUMBER({tot_nonmd})")
    ws["B14"] = (
        "Steps 4–6: Total HAC = equal-weighted mean of available z-scores (Data sheet). Cutoff = 75th percentile, non-Maryland."
    )
    ws["B14"].font = SECTION_FONT
    ws["B15"], ws["D15"] = "Cutoff in use", f'=IF($C$5="{SRC_FORMULA}",D16,D17)'
    ws["D15"].font = Font(bold=True)
    ws["B16"] = "Cutoff — formula (75th pct of Total HAC, non-MD)"
    ws["D16"] = ArrayFormula("D16", "=" + cut_f)
    ws["B17"], ws["D17"] = "Cutoff — official (pipeline)", _nan(p.get("cutoff"))
    ws["D17"].font = PY_FONT
    ws["E17"] = p.get("cutoff_method", "")
    ws["E17"].font = SUB_FONT
    ws["B18"], ws["D18"] = "Hospitals scored (excl. MD)", f"=COUNT({tot_nonmd})"
    ws["B19"], ws["D19"] = "Hospitals above cutoff in use (excl. MD)", f'=COUNTIF({tot_nonmd},">"&D15)'
    ws["B20"], ws["D20"] = "Share above cutoff", '=IFERROR(D19/D18,"—")'
    for r in (15, 16, 17):
        ws[f"D{r}"].number_format = "0.000000"
    ws["D18"].number_format = ws["D19"].number_format = F_INT
    ws["D20"].number_format = F_PCT

    # ---- What-if for the hospital selected on Hospital Lookup ----
    first = WHATIF_FIRST
    lastr = first + len(HAC_MEASURES) - 1
    ws[f"B{first - 4}"] = "What-if: selected hospital"
    ws[f"B{first - 4}"].font = SECTION_FONT
    ws[f"B{first - 3}"] = "Hospital (change on 'Hospital Lookup')"
    ws[f"D{first - 3}"] = f"=IFERROR(INDEX({b.rng('label')},{b.sel_row}),\"—\")"
    ws.merge_cells(f"D{first - 3}:H{first - 3}")
    _header(
        ws,
        first - 1,
        2,
        [
            "Measure",
            "Current value",
            "Predicted infections",
            "What-if change in value (Δ)",
            "What-if infections prevented",
            "New value",
            "Winsorized",
            "New z",
            "Weight",
            "Weighted contribution",
            "Current z (formula)",
        ],
    )
    for i, m in enumerate(HAC_MEASURES):
        r = first + i
        ws.cell(row=r, column=2, value=MEASURE_LABEL[m])
        ws.cell(row=r, column=3, value="=" + b.lookup_num(f"{m}_input"))
        if m in HAI_MEASURES and b.has(f"{m}_predicted"):
            ws.cell(row=r, column=4, value="=" + b.lookup_num(f"{m}_predicted"))
        else:
            ws.cell(row=r, column=4, value="n/a")
        _input(ws.cell(row=r, column=5, value=0))
        if m in HAI_MEASURES:
            _input(ws.cell(row=r, column=6, value=0))
        else:
            ws.cell(row=r, column=6, value="n/a")
        # SIR = observed ÷ predicted, so preventing k infections lowers SIR by k ÷ predicted. Floored at 0.
        prevented = f"IF(AND(ISNUMBER(F{r}),ISNUMBER(D{r}),N(D{r})>0),F{r}/D{r},0)"
        ws.cell(row=r, column=7, value=f'=IF(ISNUMBER(C{r}),MAX(0,C{r}+E{r}-{prevented}),"")')
        maxz = b.lookup_num(f"{m}_maxz").replace("ISNUMBER(", "ISTEXT(") if b.has(f"{m}_maxz") else '""'
        ws.cell(
            row=r,
            column=8,
            value=(
                f'=IF(ISNUMBER(G{r}),MIN(MAX(G{r},{b.param(m, "lo")}),{b.param(m, "hi")}),'
                f'IF({maxz}="Yes",{b.param(m, "hi")},""))'
            ),
        )
        ws.cell(row=r, column=9, value=f'=IF(ISNUMBER(H{r}),(H{r}-{b.param(m, "mean")})/{b.param(m, "sd")},"")')
        ws.cell(row=r, column=10, value=f"=IF(ISNUMBER(I{r}),1/COUNT($I${first}:$I${lastr}),0)")
        ws.cell(row=r, column=11, value=f'=IF(ISNUMBER(I{r}),I{r}*J{r},"")')
        ws.cell(row=r, column=12, value="=" + b.lookup_num(f"{m}_zcalc"))
        for col, fmt in zip(
            range(3, 13), [F_DEC4, F_DEC3, F_DEC4, F_INT, F_DEC4, F_DEC4, F_DEC4, "0.000", F_DEC4, F_DEC4], strict=False
        ):
            ws.cell(row=r, column=col).number_format = fmt
            ws.cell(row=r, column=col).border = BOX
        ws.cell(row=r, column=2).border = BOX
    rt = lastr + 2
    ws[f"B{rt}"], ws[f"D{rt}"] = "Current Total HAC (formula)", "=" + b.lookup_num("total_calc")
    ws[f"B{rt + 1}"], ws[f"D{rt + 1}"] = (
        "What-if Total HAC",
        f'=IF(COUNT(K{first}:K{lastr})=0,"",SUM(K{first}:K{lastr}))',
    )
    ws[f"B{rt + 2}"], ws[f"D{rt + 2}"] = "Cutoff (held constant)", f"={b.cutoff_cell}"
    ws[f"B{rt + 3}"], ws[f"D{rt + 3}"] = (
        "What-if gap (Total − cutoff)",
        f'=IF(ISNUMBER(D{rt + 1}),D{rt + 1}-D{rt + 2},"")',
    )
    ws[f"B{rt + 4}"] = "What-if status"
    ws[f"D{rt + 4}"] = (
        f'=IF(NOT(ISNUMBER(D{rt + 1})),"No HAC score",'
        f'IF(D{rt + 1}>D{rt + 2},"Above cutoff — 1% payment reduction","At or below cutoff"))'
    )
    for r in range(rt, rt + 4):
        ws[f"D{r}"].number_format = F_DEC4
        ws[f"D{r}"].font = Font(bold=True)
    tot = f"D{rt + 1}"
    ws.conditional_formatting.add(
        tot, FormulaRule(formula=[f"AND(ISNUMBER({tot}),{tot}>$D${rt + 2})"], fill=RED_FILL, font=RED_FONT)
    )
    ws.conditional_formatting.add(
        tot, FormulaRule(formula=[f"AND(ISNUMBER({tot}),{tot}<=$D${rt + 2})"], fill=GREEN_FILL, font=GREEN_FONT)
    )
    ws.conditional_formatting.add(
        f"D{rt + 4}", FormulaRule(formula=[f'LEFT(D{rt + 4},5)="Above"'], fill=RED_FILL, font=RED_FONT)
    )
    ws.conditional_formatting.add(
        f"D{rt + 4}", FormulaRule(formula=[f'LEFT(D{rt + 4},5)="At or"'], fill=GREEN_FILL, font=GREEN_FONT)
    )
    ws[f"B{rt + 6}"] = (
        "Approximate: holds every other hospital and the cutoff constant. Infections prevented convert to "
        "SIR via the predicted count from the current HAI file, which covers a later period than FY2026."
    )
    ws[f"B{rt + 6}"].font = SUB_FONT
    widths = {
        "B": 40,
        "C": 13,
        "D": 15,
        "E": 15,
        "F": 15,
        "G": 15,
        "H": 13,
        "I": 12,
        "J": 12,
        "K": 12,
        "L": 12,
        "M": 12,
        "N": 12,
        "O": 12,
        "P": 12,
    }
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.freeze_panes = "A7"


def build_lookup(b: Book) -> None:
    ws = b.wb.create_sheet("Hospital Lookup", 1)
    _title(
        ws,
        "Hospital Lookup",
        "Pick a hospital in the blue cell. Every metric below is a live lookup into the Data sheet.",
    )
    df = b.df
    default = df.loc[df.get("payment_reduction").map(_yes_no) == "Yes"] if "payment_reduction" in df else df.iloc[0:0]
    default = default.iloc[0] if len(default) else df.iloc[0]
    ws["B5"] = "Select hospital ▸"
    ws["B5"].font = Font(bold=True)
    ws["C5"] = f"{default['ccn']} — {default.get('name') or ''}"
    _input(ws["C5"])
    ws.merge_cells("C5:F5")
    dv = DataValidation(
        type="list",
        formula1=f"={b.rng('label')}",
        allow_blank=False,
        showErrorMessage=True,
        errorTitle="Pick from the list",
        error="Choose a hospital from the dropdown (CCN — Name).",
    )
    dv.add("C5")
    ws.add_data_validation(dv)
    ws["B6"], ws["C6"] = "CCN", "=LEFT(C5,6)"
    ws["B7"], ws["C7"] = "Row in Data (helper)", f"=MATCH(C6,{b.rng('ccn')},0)"
    ws["B7"].font = ws["C7"].font = HELPER_FONT

    _header(ws, 9, 2, ["Metric", "This hospital", "Peer-group average", "National average", "Source / period"])
    rows: list[tuple[str, str | None, str | None, str]] = [
        ("Hospital name (XLOOKUP)", "name", None, "Hospital General Information (xubh-q36u)"),
        ("City", "city", None, "xubh-q36u"),
        ("State", "state", None, "xubh-q36u"),
        ("Hospital type", "hospital_type", None, "xubh-q36u"),
        ("Peer group", "peer_label", None, "Cost report + rules in config/peer_rules.yaml"),
        ("Peer group key", "peer_key", None, "Derived"),
        ("Beds", "beds", F_INT, "Hospital Provider Cost Report"),
        ("Total HAC Score (official)", "hac_total", F_DEC4, "HAC Reduction Program FY2026 (yq43-i98g)"),
        ("Total HAC Score (rebuilt, formula)", "total_calc", F_DEC4, "Recomputed on Data sheet"),
        ("Gap to cutoff (+ = above)", "hac_gap", F_DEC4, "yq43-i98g; cutoff from HAC Calculator"),
        ("HAC status", "hac_status", None, "Derived (see Definitions)"),
        ("Payment reduction (official)", "payment_reduction", None, "yq43-i98g"),
        ("HAC national percentile (higher = worse)", "hac_percentile", F_PCT_PTS, "Derived"),
    ]
    for m in HAC_MEASURES:
        rows.append(
            (
                f"{MEASURE_LABEL[m]} — value",
                f"{m}_input",
                F_DEC4 if m == "psi90" else F_DEC3,
                "yq43-i98g (PSI 90: 07/2022–06/2024; HAI: 01/2023–12/2024)",
            )
        )
        rows.append((f"{MEASURE_LABEL[m]} — z (official)", f"{m}_z_official", F_DEC4, "yq43-i98g"))
    rows += [
        ("Tipping measure", "tipping_measure", None, "Patients-to-prevent (approximate)"),
        ("Events to prevent at tipping measure", "tipping_events", F_INT, "Patients-to-prevent (approximate)"),
        ("Trajectory Total HAC (proxy)", "traj_total", F_DEC4, "Current HAI file (77hc-ibv8), FY2026 distribution"),
        ("Nurse communication 'Always' %", "nurse_comm_top_box", F_PCT_PTS, "HCAHPS (dgck-syfz) H_COMP_1_A_P"),
        ("Discharge information 'Yes' %", "discharge_info_yes", F_PCT_PTS, "HCAHPS (dgck-syfz) H_COMP_6_Y_P"),
        ("HCAHPS summary star rating", "hcahps_star", "0", "dgck-syfz"),
        ("ERR — Heart failure", "err_hf", F_DEC4, "HRRP FY2026 (9n3s-kdb3)"),
        ("ERR — Pneumonia", "err_pn", F_DEC4, "9n3s-kdb3"),
        ("ERR — AMI", "err_ami", F_DEC4, "9n3s-kdb3"),
        ("ERR — COPD", "err_copd", F_DEC4, "9n3s-kdb3"),
        ("ERR — CABG", "err_cabg", F_DEC4, "9n3s-kdb3"),
        ("ERR — Hip/knee", "err_hip_knee", F_DEC4, "9n3s-kdb3"),
        ("Est. 1% HAC exposure ($)", "exposure_usd", F_USD, "Medicare Inpatient Hospitals – by Provider (estimate)"),
        ("VBP Total Performance Score", "vbp_tps", F_DEC2, "HVBP TPS (ypbt-wvdk)"),
    ]
    peer_row = None
    r = 10
    cells: dict[str, str] = {}
    for label, key, fmt, src in rows:
        if key is None or not b.has(key):
            continue
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=3, value=b.lookup(key))
        cells[key] = f"C{r}"
        if key == "peer_key":
            peer_row = r
        numeric = fmt not in (None, "0") or key == "hcahps_star"
        if numeric and key not in ("hac_status", "tipping_events"):
            ws.cell(row=r, column=4, value=None)  # filled after peer_row is known
            ws.cell(row=r, column=5, value=f'=IFERROR(AVERAGE({b.rng(key)}),"—")')
        ws.cell(row=r, column=6, value=src).font = SUB_FONT
        for col in (3, 4, 5):
            ws.cell(row=r, column=col).number_format = fmt or "General"
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="right")
        for col in range(2, 7):
            ws.cell(row=r, column=col).border = BOX
        r += 1
    # Peer averages via AVERAGEIFS on peer key.
    if peer_row and b.has("peer_key"):
        for rr in range(10, r):
            if ws.cell(row=rr, column=5).value:  # numeric row
                key = next(k for k, a in cells.items() if a == f"C{rr}")
                ws.cell(
                    row=rr, column=4, value=f'=IFERROR(AVERAGEIFS({b.rng(key)},{b.rng("peer_key")},$C${peer_row}),"—")'
                )
    # Conditional formatting vs the cutoff.
    for key in ("hac_total", "total_calc"):
        if key in cells:
            a = cells[key]
            ws.conditional_formatting.add(
                a, FormulaRule(formula=[f"AND(ISNUMBER({a}),{a}>{b.cutoff_cell})"], fill=RED_FILL, font=RED_FONT)
            )
            ws.conditional_formatting.add(
                a, FormulaRule(formula=[f"AND(ISNUMBER({a}),{a}<={b.cutoff_cell})"], fill=GREEN_FILL, font=GREEN_FONT)
            )
    if "hac_gap" in cells:
        a = cells["hac_gap"]
        ws.conditional_formatting.add(
            a, CellIsRule(operator="greaterThan", formula=["0"], fill=RED_FILL, font=RED_FONT)
        )
    if "payment_reduction" in cells:
        a = cells["payment_reduction"]
        ws.conditional_formatting.add(a, FormulaRule(formula=[f'{a}="Yes"'], fill=RED_FILL, font=RED_FONT))
    for k in ("err_hf", "err_pn", "err_ami", "err_copd", "err_cabg", "err_hip_knee"):
        if k in cells:
            a = cells[k]
            ws.conditional_formatting.add(a, FormulaRule(formula=[f"AND(ISNUMBER({a}),{a}>1)"], fill=AMBER_FILL))
    ws[f"B{r + 1}"] = (
        "Cutoff used for colouring: 'HAC Calculator'!D15. Lookups use XLOOKUP with an INDEX/MATCH fallback "
        "for older Excel. '—' means CMS did not report the value."
    )
    ws[f"B{r + 1}"].font = SUB_FONT
    for k, v in {"B": 44, "C": 22, "D": 18, "E": 18, "F": 52}.items():
        ws.column_dimensions[k].width = v
    b.lookup_cells = cells


def build_ptp(b: Book) -> None:
    ws = b.wb.create_sheet("Patients to Prevent")
    _title(
        ws,
        "Patients to Prevent — formula version of the CareSignal method",
        "For the hospital selected on 'Hospital Lookup'. Approximate: peers and the cutoff are held constant.",
    )
    ws["B5"], ws["D5"] = "Hospital", f"=IFERROR(INDEX({b.rng('label')},{b.sel_row}),\"—\")"
    ws.merge_cells("D5:H5")
    # Calibrated mode uses the official CMS total (what the site uses); formula mode uses the live rebuild.
    off = b.lookup_num("hac_total") if b.has("hac_total") else '""'
    ws["B6"] = "Total HAC Score"
    ws["D6"] = (
        f"=IF('HAC Calculator'!$C$5=\"{SRC_FORMULA}\",{b.lookup_num('total_calc')},"
        f"IF(ISNUMBER({off}),{off},{b.lookup_num('total_calc')}))"
    )
    ws["B7"], ws["D7"] = "Cutoff", f"={b.cutoff_cell}"
    ws["B8"], ws["D8"] = "Gap Δ = Total − cutoff", '=IF(ISNUMBER(D6),D6-D7,"")'
    ws["B9"], ws["D9"] = "Measures with a result (n)", "=COUNT(F13:F18)"
    ws["B10"], ws["D10"] = "Weight per measure (1/n)", '=IF(D9>0,1/D9,"")'
    for r in (6, 7, 8, 10):
        ws[f"D{r}"].number_format = F_DEC4
    ws.conditional_formatting.add("D8", FormulaRule(formula=["AND(ISNUMBER(D8),D8>0)"], fill=RED_FILL, font=RED_FONT))
    ws.conditional_formatting.add(
        "D8", FormulaRule(formula=["AND(ISNUMBER(D8),D8<=0)"], fill=GREEN_FILL, font=GREEN_FONT)
    )

    _header(
        ws,
        12,
        2,
        [
            "Measure",
            "Current value (SIR / rate)",
            "Predicted infections",
            "Winsor low",
            "Current z",
            "Required Δz = Δ / weight",
            "Target winsorized value",
            "Required value change",
            "Infections to prevent (formula)",
            "Python events",
            "Note",
        ],
    )
    for i, m in enumerate(HAC_MEASURES):
        r = 13 + i
        ws.cell(row=r, column=2, value=MEASURE_LABEL[m])
        ws.cell(row=r, column=3, value="=" + b.lookup_num(f"{m}_input"))
        ws.cell(
            row=r,
            column=4,
            value=("=" + b.lookup_num(f"{m}_predicted")) if m in HAI_MEASURES and b.has(f"{m}_predicted") else "n/a",
        )
        ws.cell(row=r, column=5, value=f"={b.param(m, 'lo')}")
        ws.cell(row=r, column=6, value="=" + b.lookup_num(f"{m}_zcalc"))
        ws.cell(row=r, column=7, value=f'=IF(AND(ISNUMBER(F{r}),ISNUMBER($D$8),$D$8>0),$D$8/$D$10,"")')
        # target winsorized value = current winsorized − Δz × SD
        w_now = f"MIN(MAX(C{r},{b.param(m, 'lo')}),{b.param(m, 'hi')})"
        ws.cell(row=r, column=8, value=f'=IF(ISNUMBER(G{r}),{w_now}-G{r}*{b.param(m, "sd")},"")')
        ws.cell(row=r, column=9, value=f'=IF(ISNUMBER(H{r}),IF(H{r}<E{r}-1E-12,"Not reachable",C{r}-H{r}),"")')
        if m in HAI_MEASURES:
            sd, lo, hi = b.param(m, "sd"), b.param(m, "lo"), b.param(m, "hi")
            let = (
                f"_xlfn.LET(_xlpm.dz,$D$8/$D$10,_xlpm.w,MIN(MAX(C{r},{lo}),{hi}),_xlpm.t,_xlpm.w-_xlpm.dz*{sd},"
                f'IF(_xlpm.t<{lo}-1E-12,"Not reachable",CEILING((C{r}-_xlpm.t)*D{r}-1E-9,1)))'
            )
            plain = f'IF(H{r}<E{r}-1E-12,"Not reachable",CEILING((C{r}-H{r})*D{r}-1E-9,1))'
            ws.cell(
                row=r,
                column=10,
                value=(
                    f'=IF(NOT(AND(ISNUMBER($D$8),ISNUMBER(C{r}),ISNUMBER(D{r}))),"—",'
                    f'IF($D$8<=0,"Already clear",IFERROR({let},{plain})))'
                ),
            )
            note = "Events = Δ SIR × predicted (SIR = observed ÷ predicted)"
        else:
            ws.cell(row=r, column=10, value="n/a")
            note = "Composite rate: shows required change in PSI 90 value, not an event count"
        py = ws.cell(
            row=r, column=11, value=("=" + b.lookup_num(f"ptp_{m}_events")) if b.has(f"ptp_{m}_events") else None
        )
        py.font = PY_FONT
        ws.cell(row=r, column=12, value=note).font = SUB_FONT
        for col, fmt in zip(
            range(3, 12), [F_DEC4, F_DEC3, F_DEC4, F_DEC4, F_DEC4, F_DEC4, F_DEC4, F_INT, F_INT], strict=False
        ):
            ws.cell(row=r, column=col).number_format = fmt
            ws.cell(row=r, column=col).border = BOX
        ws.cell(row=r, column=2).border = BOX
    ws["B20"], ws["D20"] = (
        "Tipping measure (fewest events, formula)",
        ('=IFERROR(INDEX(B13:B18,MATCH(MIN(J13:J18),J13:J18,0)),"—")'),
    )
    ws["B21"], ws["D21"] = "Events at tipping measure", '=IF(COUNT(J13:J18)=0,"—",MIN(J13:J18))'
    ws["D21"].number_format = F_INT
    ws["B23"] = (
        "Method: Δ = Total − cutoff; Δz = Δ ÷ weight; target winsorized value = current winsorized − Δz × SD; "
        "infections = ⌈(current SIR − target) × predicted⌉. If the target falls below the winsor floor the measure "
        "alone cannot clear the line ('Not reachable'). Predicted counts come from the current HAI file (proxy period)."
    )
    ws["B23"].alignment = Alignment(wrap_text=True)
    ws["B23"].font = SUB_FONT
    ws.merge_cells("B23:L23")
    ws.row_dimensions[23].height = 46
    for k, v in {
        "B": 38,
        "C": 13,
        "D": 13,
        "E": 11,
        "F": 11,
        "G": 13,
        "H": 13,
        "I": 13,
        "J": 15,
        "K": 10,
        "L": 52,
    }.items():
        ws.column_dimensions[k].width = v


def _peer_array(b: Book, key: str, peer: str, fn: str, x: str | None = None) -> str:
    cond = f"({b.rng('peer_key')}={peer})*ISNUMBER({b.rng(key)})"
    arr = f'IF({cond},{b.rng(key)},"")'
    return f"{fn}({arr},{x})" if x else f"{fn}({arr})"


def build_peer(b: Book) -> None:
    ws = b.wb.create_sheet("Peer Benchmark")
    _title(
        ws,
        "Peer Benchmark",
        "Percentile rank (PERCENTRANK.INC) within the selected hospital's peer group and nationally.",
    )
    ws["B5"], ws["D5"] = "Hospital", f"=IFERROR(INDEX({b.rng('label')},{b.sel_row}),\"—\")"
    ws.merge_cells("D5:H5")
    ws["B6"], ws["D6"] = (
        "Peer group",
        f"=IFERROR(INDEX({b.rng('peer_label')},{b.sel_row}),\"—\")" if b.has("peer_label") else "—",
    )
    ws.merge_cells("D6:H6")
    ws["B7"], ws["D7"] = (
        "Peer key",
        f"=IFERROR(INDEX({b.rng('peer_key')},{b.sel_row}),\"\")" if b.has("peer_key") else "",
    )
    ws["B7"].font = ws["D7"].font = HELPER_FONT
    _header(
        ws,
        9,
        2,
        [
            "Metric",
            "Direction",
            "This hospital",
            "Peers with value",
            "Peer median",
            "Peer percentile rank",
            "Better than % of peers",
            "National percentile rank",
            "Better than % nationally",
        ],
    )
    metrics = [("Total HAC Score (official)", "hac_total", "Lower is better", F_DEC4)]
    metrics += [
        (MEASURE_LABEL[m], f"{m}_input", "Lower is better", F_DEC4 if m == "psi90" else F_DEC3) for m in HAC_MEASURES
    ]
    metrics += [
        ("Nurse communication 'Always' %", "nurse_comm_top_box", "Higher is better", F_PCT_PTS),
        ("Discharge information 'Yes' %", "discharge_info_yes", "Higher is better", F_PCT_PTS),
        ("HCAHPS summary star", "hcahps_star", "Higher is better", "0.0"),
        ("ERR — Heart failure", "err_hf", "Lower is better", F_DEC4),
        ("ERR — Pneumonia", "err_pn", "Lower is better", F_DEC4),
        ("ERR — AMI", "err_ami", "Lower is better", F_DEC4),
        ("ERR — COPD", "err_copd", "Lower is better", F_DEC4),
    ]
    r = 10
    for label, key, direction, fmt in metrics:
        if not b.has(key) or not b.has("peer_key"):
            continue
        x = f"D{r}"
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=3, value=direction)
        ws.cell(row=r, column=4, value="=" + b.lookup_num(key)).number_format = fmt
        ws.cell(
            row=r, column=5, value=f'=IF($D$7="","—",COUNTIFS({b.rng("peer_key")},$D$7,{b.rng(key)},"<>"))'
        ).number_format = F_INT
        ref = f"E{r}:E{r}"
        ws[f"F{r}"] = ArrayFormula(f"F{r}", f'=IFERROR(IF($D$7="","—",{_peer_array(b, key, "$D$7", "MEDIAN")}),"—")')
        ws[f"F{r}"].number_format = fmt
        ws[f"G{r}"] = ArrayFormula(
            f"G{r}",
            f'=IFERROR(IF(OR($D$7="",NOT(ISNUMBER({x}))),"—",'
            f'{_peer_array(b, key, "$D$7", "_xlfn.PERCENTRANK.INC", x)}),"—")',
        )
        ws[f"G{r}"].number_format = F_PCT
        lower = direction.startswith("Lower")
        ws[f"H{r}"] = f'=IF(ISNUMBER(G{r}),{"1-G" if lower else "G"}{r},"—")'
        ws[f"H{r}"].number_format = F_PCT
        ws[f"I{r}"] = f'=IF(ISNUMBER({x}),IFERROR(_xlfn.PERCENTRANK.INC({b.rng(key)},{x}),"—"),"—")'
        ws[f"I{r}"].number_format = F_PCT
        ws[f"J{r}"] = f'=IF(ISNUMBER(I{r}),{"1-I" if lower else "I"}{r},"—")'
        ws[f"J{r}"].number_format = F_PCT
        for col in range(2, 11):
            ws.cell(row=r, column=col).border = BOX
        ws.conditional_formatting.add(
            f"H{r}", FormulaRule(formula=[f"AND(ISNUMBER(H{r}),H{r}<0.25)"], fill=RED_FILL, font=RED_FONT)
        )
        ws.conditional_formatting.add(
            f"H{r}", FormulaRule(formula=[f"AND(ISNUMBER(H{r}),H{r}>=0.75)"], fill=GREEN_FILL, font=GREEN_FONT)
        )
        _ = ref
        r += 1
    ws[f"B{r + 1}"] = (
        "PERCENTRANK.INC returns the share of peers with a lower value. For lower-is-better measures, "
        "'Better than % of peers' = 1 − rank. Ties and small peer groups make ranks coarse."
    )
    ws[f"B{r + 1}"].font = SUB_FONT
    for k, v in {"B": 38, "C": 15, "D": 13, "E": 11, "F": 12, "G": 13, "H": 13, "I": 13, "J": 13}.items():
        ws.column_dimensions[k].width = v


def build_hcahps(b: Book) -> None:
    ws = b.wb.create_sheet("HCAHPS")
    _title(
        ws,
        "HCAHPS — nurse communication and discharge information",
        "Patient survey (dgck-syfz). Top-box % vs peers, state and nation. Percentile ranks: higher = better.",
    )
    ws["B5"], ws["D5"] = "Hospital", f"=IFERROR(INDEX({b.rng('label')},{b.sel_row}),\"—\")"
    ws.merge_cells("D5:H5")
    ws["B6"], ws["D6"] = (
        "Peer key",
        f"=IFERROR(INDEX({b.rng('peer_key')},{b.sel_row}),\"\")" if b.has("peer_key") else "",
    )
    ws["B7"], ws["D7"] = "State", f"=IFERROR(INDEX({b.rng('state')},{b.sel_row}),\"\")"
    for c in ("B6", "D6", "B7", "D7"):
        ws[c].font = HELPER_FONT
    _header(
        ws,
        9,
        2,
        [
            "Measure",
            "This hospital",
            "Peer median",
            "Gap to peer median (pts)",
            "Peer percentile",
            "State median",
            "National median",
            "National percentile",
            "Python national pctl",
            "Python peer pctl",
        ],
    )
    items = [
        (
            "Nurse communication — 'Always' % (H_COMP_1_A_P)",
            "nurse_comm_top_box",
            "nurse_comm_pctl_nat",
            "nurse_comm_pctl_peer",
        ),
        ("Discharge information — 'Yes' % (H_COMP_6_Y_P)", "discharge_info_yes", None, None),
        ("Care transition % (H_COMP_7)", "care_transition", None, None),
    ]
    r = 10
    for label, key, py_nat, py_peer in items:
        if not b.has(key):
            continue
        x = f"C{r}"
        ws.cell(row=r, column=2, value=label)
        ws[f"C{r}"] = "=" + b.lookup_num(key)
        if b.has("peer_key"):
            ws[f"D{r}"] = ArrayFormula(
                f"D{r}", f'=IFERROR(IF($D$6="","—",{_peer_array(b, key, "$D$6", "MEDIAN")}),"—")'
            )
            ws[f"F{r}"] = ArrayFormula(
                f"F{r}",
                f'=IFERROR(IF(OR($D$6="",NOT(ISNUMBER({x}))),"—",'
                f'{_peer_array(b, key, "$D$6", "_xlfn.PERCENTRANK.INC", x)}),"—")',
            )
        ws[f"E{r}"] = f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),C{r}-D{r},"—")'
        st = f'IF(({b.rng("state")}=$D$7)*ISNUMBER({b.rng(key)}),{b.rng(key)},"")'
        ws[f"G{r}"] = ArrayFormula(f"G{r}", f'=IFERROR(MEDIAN({st}),"—")')
        ws[f"H{r}"] = f'=IFERROR(MEDIAN({b.rng(key)}),"—")'
        ws[f"I{r}"] = f'=IF(ISNUMBER({x}),IFERROR(_xlfn.PERCENTRANK.INC({b.rng(key)},{x}),"—"),"—")'
        if py_nat and b.has(py_nat):
            ws[f"J{r}"] = "=" + b.lookup_num(py_nat)
            ws[f"J{r}"].font = PY_FONT
        if py_peer and b.has(py_peer):
            ws[f"K{r}"] = "=" + b.lookup_num(py_peer)
            ws[f"K{r}"].font = PY_FONT
        for col, fmt in zip(
            "CDEFGHIJK",
            [F_PCT_PTS, F_PCT_PTS, "+0.0;-0.0;0.0", F_PCT, F_PCT_PTS, F_PCT_PTS, F_PCT, F_PCT_PTS, F_PCT_PTS],
            strict=False,
        ):
            ws[f"{col}{r}"].number_format = fmt
            ws[f"{col}{r}"].border = BOX
        ws[f"B{r}"].border = BOX
        ws.conditional_formatting.add(f"E{r}", FormulaRule(formula=[f"AND(ISNUMBER(E{r}),E{r}<0)"], fill=AMBER_FILL))
        r += 1
    ws[f"B{r + 1}"] = (
        "Python percentiles are on a 0–100 scale and may use a different tie method than PERCENTRANK.INC. "
        "HCAHPS period differs from the HAC program period; see README."
    )
    ws[f"B{r + 1}"].font = SUB_FONT
    for k, v in {"B": 48, "C": 12, "D": 12, "E": 13, "F": 12, "G": 12, "H": 12, "I": 12, "J": 12, "K": 12}.items():
        ws.column_dimensions[k].width = v


def build_dq(b: Book) -> None:
    ws = b.wb.create_sheet("DQ Log")
    _title(
        ws,
        "Data quality log",
        "Every validation rule the pipeline ran, with record counts (data/processed/dq_log.csv).",
    )
    dq = b.inp.dq
    if dq is None or dq.empty:
        ws["B5"] = "dq_log.csv was not available when this workbook was built."
        return
    cols = list(dq.columns)
    _header(ws, 5, 2, cols)
    for i, rec in enumerate(dq.itertuples(index=False), start=6):
        for j, v in enumerate(rec, start=2):
            val = v
            if isinstance(v, str) and v.lstrip("-").isdigit():
                val = int(v)
            c = ws.cell(row=i, column=j, value=val)
            c.border = BOX
    for j, c in enumerate(cols, start=2):
        ws.column_dimensions[get_column_letter(j)].width = 60 if "desc" in c.lower() else 16
    ws.freeze_panes = "B6"
    if any(c.lower() in ("passed", "status") for c in cols):
        pc = next(c for c in cols if c.lower() in ("passed", "status"))
        L = get_column_letter(cols.index(pc) + 2)
        rng = f"{L}6:{L}{5 + len(dq)}"
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'OR(UPPER({L}6)="FALSE",UPPER({L}6)="FAIL")'], fill=RED_FILL, font=RED_FONT)
        )


def build_definitions(b: Book) -> None:
    ws = b.wb.create_sheet("Definitions")
    _title(
        ws,
        "Definitions",
        "Measure dictionary with CareSignal's nursing category and rationale (config/measure_map.yaml).",
    )
    _header(ws, 5, 2, ["Measure ID", "Label", "Program", "Nursing category", "Direction", "Source", "Rationale"])
    r = 6
    mm = b.inp.measure_map
    if not mm:
        mm = [
            {"id": m, "label": MEASURE_LABEL[m], "program": "HAC", "category": MEASURE_CATEGORY[m]}
            for m in HAC_MEASURES
        ]
    for it in mm:
        vals = [
            it.get("id"),
            it.get("label"),
            it.get("program"),
            it.get("category"),
            it.get("direction"),
            it.get("source"),
            it.get("rationale"),
        ]
        for j, v in enumerate(vals, start=2):
            c = ws.cell(row=r, column=j, value=v if not isinstance(v, list | dict) else json.dumps(v))
            c.border = BOX
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="Terms").font = SECTION_FONT
    terms = [
        ("SIR", "Standardized Infection Ratio = observed ÷ predicted infections (NHSN). 1.0 = as predicted."),
        ("Winsorize", "Cap each measure at its 5th and 95th percentile across hospitals before standardizing."),
        ("z-score", "(winsorized value − mean) ÷ SD. Higher = worse on every HAC measure."),
        ("Total HAC Score", "Equal-weighted average of the z-scores a hospital has (1/n each)."),
        ("Cutoff", "75th percentile of Total HAC Scores (non-Maryland). Above it = 1% Medicare FFS payment reduction."),
        (
            "HAC status",
            "penalized = official payment reduction; near = clear but ≤2 added infections from crossing; clear; no_score.",
        ),
        (
            "Patients to prevent",
            "Approximate whole infections to avoid in one measure to clear the cutoff, holding peers constant.",
        ),
        ("ERR", "HRRP Excess Readmission Ratio; above 1.0 = more readmissions than expected."),
        ("Top-box", "Share of HCAHPS respondents giving the most positive answer (e.g. 'Always')."),
        (
            "Peer group",
            "Bed size × teaching intensity × urban/rural × Census region; falls back to bed × teaching if <15 hospitals.",
        ),
    ]
    for t, d in terms:
        r += 1
        ws.cell(row=r, column=2, value=t).font = Font(bold=True)
        ws.cell(row=r, column=3, value=d)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    for k, v in {"B": 16, "C": 34, "D": 12, "E": 16, "F": 14, "G": 14, "H": 70}.items():
        ws.column_dimensions[k].width = v


def build_readme(b: Book) -> None:
    ws = b.wb.active
    ws.title = "README"
    _title(ws, "CareSignal Workbook", f"Built {b.inp.built_at} from public CMS data. Hospitals: {b.n:,}.")
    lines = [
        (
            "Purpose",
            "How close is a hospital to the Medicare HAC Reduction Program penalty line, which nursing-sensitive "
            "measure drives it, and how many infections would need to be prevented to get clear? Every derived "
            "number on the calculator sheets is a live formula you can audit.",
        ),
        (
            "How to use",
            "1) Pick a hospital in the blue cell on 'Hospital Lookup'. 2) 'HAC Calculator' rebuilds the FY2026 "
            "score and lets you enter what-if changes (blue cells). 3) 'Patients to Prevent', 'Peer Benchmark' "
            "and 'HCAHPS' all follow the selected hospital.",
        ),
        (
            "Colour key",
            "Blue cells = inputs. Grey 'Python' values = pipeline outputs for audit. Red = above cutoff / worse "
            "than benchmark. '—' = not reported by CMS (never treated as zero).",
        ),
        (
            "Compatibility",
            "Lookups use XLOOKUP with an INDEX/MATCH fallback; one formula uses LET with a plain fallback. "
            "Works in Excel 2010+ (best in Excel 365) and LibreOffice 24.8+. Peer medians/percentiles are array formulas.",
        ),
        (
            "Limitations",
            "Patients-to-prevent holds peers and the cutoff constant; the cutoff moves every year. Predicted counts "
            "come from the current HAI file, a later period than FY2026. Dollar exposure is an order-of-magnitude "
            "estimate (1% of Medicare inpatient payments).",
        ),
    ]
    r = 5
    for k, v in lines:
        ws.cell(row=r, column=2, value=k).font = Font(bold=True, color=TEAL)
        c = ws.cell(row=r, column=3, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
        ws.row_dimensions[r].height = 48
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="Sources").font = SECTION_FONT
    r += 1
    _header(ws, r, 2, ["Dataset ID", "Title", "CMS modified", "Rows", "Role"])
    for s in b.inp.sources:
        r += 1
        vals = [s.get("id"), s.get("title"), s.get("modified"), s.get("rows"), s.get("role")]
        for j, v in enumerate(vals, start=2):
            c = ws.cell(row=r, column=j, value=_nan(v) if not isinstance(v, dict | list) else json.dumps(v))
            c.border = BOX
        ws.cell(row=r, column=5).number_format = F_INT
    if not b.inp.sources:
        r += 1
        ws.cell(row=r, column=2, value="Source manifest not available at build time; see the site's Data Quality page.")
    r += 2
    ws.cell(row=r, column=2, value="Refresh").font = Font(bold=True, color=TEAL)
    ws.cell(row=r, column=3, value=f"Regenerated weekly by GitHub Actions. This copy: {b.inp.built_at}.")
    for k, v in {"B": 18, "C": 48, "D": 14, "E": 10, "F": 48, "G": 10}.items():
        ws.column_dimensions[k].width = v


def build(inp: Inputs | None = None, out: Path = OUT) -> Path:
    inp = inp or load_inputs()
    b = Book(inp)
    build_readme(b)
    build_data(b)
    build_calculator(b)
    build_lookup(b)
    build_ptp(b)
    build_peer(b)
    build_hcahps(b)
    build_dq(b)
    build_definitions(b)
    order = [
        "README",
        "Hospital Lookup",
        "HAC Calculator",
        "Patients to Prevent",
        "Peer Benchmark",
        "HCAHPS",
        "Data",
        "DQ Log",
        "Definitions",
    ]
    b.wb._sheets = [b.wb[n] for n in order]
    b.wb.active = 1
    b.wb.calculation.fullCalcOnLoad = True  # no cached values: force Excel/LibreOffice to compute on open
    out.parent.mkdir(parents=True, exist_ok=True)
    b.wb.save(out)
    return out


def main() -> None:
    out = build()
    print(f"export_excel: wrote {out.relative_to(ROOT)} ({out.stat().st_size / 1e6:.1f} MB)")


if __name__ == "__main__":
    main()
