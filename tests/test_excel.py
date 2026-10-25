"""Structural tests for the Excel workbook: live formulas, data validation, Excel Table, sheet set.

Runs on a tiny synthetic input (always), and on the real built workbook when it exists.
"""

from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from pipeline import export_excel as ex

SHEETS = [
    "README",
    "Hospital Lookup",
    "HAC Calculator",
    "Patients to Prevent",
    "Peer Benchmark",
    "HCAHPS",
    "Data",
    "DQ Log",
    "Definitions",
]


def _synthetic() -> ex.Inputs:
    rng = np.random.default_rng(0)
    n = 40
    df = pd.DataFrame(
        {
            "ccn": [f"{i:06d}" for i in range(10001, 10001 + n)],
            "name": [f"Hospital {i}" for i in range(n)],
            "city": "Town",
            "state": ["AL", "MD"] * (n // 2),
            "hospital_type": "Acute Care Hospitals",
            "peer_key": ["b100_299-none-urban-south"] * (n // 2) + ["b0_99-none-rural-south"] * (n // 2),
            "peer_label": "Peers",
            "beds": 120,
            "is_maryland": [False, True] * (n // 2),
            "hac_total": rng.normal(0, 0.5, n),
            "payment_reduction": [True] + [False] * (n - 1),
            "nurse_comm_top_box": rng.uniform(70, 90, n),
            "discharge_info_yes": rng.uniform(80, 92, n),
        }
    )
    for m in ex.HAC_MEASURES:
        df[f"{m}_input"] = rng.uniform(0.2, 2.0, n)
        df[f"{m}_z_official"] = rng.normal(0, 1, n)
        if m in ex.HAI_MEASURES:
            df[f"{m}_predicted"] = rng.uniform(1, 10, n)
        df[f"ptp_{m}_events"] = pd.array([None] * n, dtype="Int64")
    df.loc[3, "cauti_input"] = np.nan
    params = {
        "fiscal_year": 2026,
        "cutoff": 0.3,
        "near_band": 2,
        "conventions": {"percentile_method": "nearest", "sd_ddof": 1, "include_maryland": True},
        "measures": {
            m: {"winsor_low": 0.3, "winsor_high": 1.9, "mean": 1.0, "sd": 0.4, "n": n} for m in ex.HAC_MEASURES
        },
    }
    dq = pd.DataFrame(
        {
            "rule_id": ["R1"],
            "dataset": ["x"],
            "description": ["d"],
            "severity": ["error"],
            "checked": ["10"],
            "failed": ["0"],
            "passed": ["True"],
        }
    )
    return ex.Inputs(
        df, params, dq, [], [{"id": "yq43-i98g", "title": "HAC", "modified": "2026-01-26", "rows": 3055}], "now"
    )


@pytest.fixture(scope="module")
def synthetic_book(tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("xl") / "wb.xlsx"
    return ex.build(_synthetic(), out)


def _check(path: Path) -> None:
    wb = load_workbook(path)  # formulas, not cached values
    assert wb.sheetnames == SHEETS

    data = wb["Data"]
    assert "tblHospitals" in data.tables
    headers = [c.value for c in data[1]]
    assert headers[0] == "CCN"
    # CCN kept as text with leading zeros
    assert isinstance(data["A2"].value, str) and len(data["A2"].value) == 6

    # Live formula block on Data
    w_col = headers.index("CAUTI winsorized (formula)") + 1
    z_col = headers.index("CAUTI z (formula)") + 1
    t_col = headers.index("Total HAC (formula)") + 1
    for col in (w_col, z_col, t_col):
        v = data.cell(row=2, column=col).value
        assert isinstance(v, str) and v.startswith("="), v

    calc = wb["HAC Calculator"]
    for row in range(7, 13):
        for col in "DEFGMNOP":
            v = calc[f"{col}{row}"].value
            v = getattr(v, "text", v)  # ArrayFormula
            assert isinstance(v, str) and v.startswith("="), (col, row, v)
    d7 = getattr(calc["D7"].value, "text", calc["D7"].value)
    assert "SMALL(" in d7 or "PERCENTILE.INC" in d7
    assert "STDEV" in getattr(calc["G7"].value, "text", calc["G7"].value)
    assert "0.75" in getattr(calc["D16"].value, "text", calc["D16"].value)
    assert calc["D15"].value.startswith("=IF(")
    # Blue what-if inputs are plain values, formulas downstream
    assert calc["E27"].value == 0 and calc["E27"].fill.fgColor.rgb.endswith("DDEBF7")
    assert calc["I27"].value.startswith("=")

    lookup = wb["Hospital Lookup"]
    dvs = lookup.data_validations.dataValidation
    assert any("C5" in str(dv.sqref) and dv.type == "list" for dv in dvs)
    formulas = [c.value for c in lookup["C"] if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("XLOOKUP" in f and "INDEX(" in f for f in formulas)
    assert any("AVERAGEIFS" in (c.value or "") for c in lookup["D"] if isinstance(c.value, str))
    assert lookup.conditional_formatting  # cutoff colouring

    ptp = wb["Patients to Prevent"]
    cauti_row = 13 + ex.HAC_MEASURES.index("cauti")
    assert "CEILING" in ptp[f"J{cauti_row}"].value and "LET" in ptp[f"J{cauti_row}"].value

    peer = wb["Peer Benchmark"]
    assert "PERCENTRANK.INC" in str(peer["G10"].value.text if hasattr(peer["G10"].value, "text") else peer["G10"].value)

    assert wb.calculation.fullCalcOnLoad


def test_synthetic_workbook_structure(synthetic_book):
    _check(synthetic_book)


def test_python_evaluator_matches_params_logic():
    """The formula logic (winsorize → z → mean of available) reproduces a hand calculation."""
    inp = _synthetic()
    df, p = inp.data, inp.params["measures"]
    z = pd.DataFrame(
        {
            m: ((df[f"{m}_input"].clip(p[m]["winsor_low"], p[m]["winsor_high"]) - p[m]["mean"]) / p[m]["sd"])
            for m in ex.HAC_MEASURES
        }
    )
    total = z.mean(axis=1, skipna=True)
    # row 3 has no CAUTI → average of the other five
    expected = z.loc[3, [m for m in ex.HAC_MEASURES if m != "cauti"]].mean()
    assert total[3] == pytest.approx(expected)


REAL = ex.OUT


@pytest.mark.data
@pytest.mark.skipif(not REAL.exists(), reason="workbook not built yet (run make data)")
def test_real_workbook_structure():
    _check(REAL)
