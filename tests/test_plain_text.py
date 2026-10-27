"""Plain language contract (docs/STYLE_GUIDE.md): no visible text on the site or in the workbook carries an
em dash, en dash, or a hyphen used as punctuation or in a compound word. Identifiers (CCNs, peer keys,
dataset ids, ISO dates, URLs, file names) are allowed to keep their hyphens."""

from __future__ import annotations

import json

import pytest

from pipeline.common import WEB_DATA, WEB_DOWNLOADS
from pipeline.plain_text import ID_KEYS, offending, plain

WORKBOOK = WEB_DOWNLOADS / "CareSignal_Workbook.xlsx"


def test_plain_rewrites_dashes():
    assert plain("SSI – colon surgery") == "SSI, colon surgery"
    assert plain("Voluntary non-profit - Private") == "Voluntary nonprofit, Private"
    assert plain("100–299 beds") == "100 to 299 beds"
    assert plain("Nursing-driven") == "Nursing driven"
    assert plain("—") == "N/A"
    assert plain("r = -0.07") == "r = −0.07"


def test_plain_keeps_identifiers():
    for s in ("yq43-i98g", "b100_299-minor-urban-south", "2024-06-30", "https://data.cms.gov/a-b", "READM-30-HF-HRRP"):
        assert plain(s) == s
        assert offending(s) == []


def test_offending_flags_dashes():
    assert offending("Above cutoff — 1%")
    assert offending("Steps 1–3")
    assert offending("Nursing-driven")
    assert offending("CENTER- COLLEGE")
    assert offending("Total −0.25") == []


def _walk(obj, key=None, path=""):
    if isinstance(obj, dict):
        for k, v in obj.items():
            yield from _walk(v, k, f"{path}.{k}")
    elif isinstance(obj, list):
        for v in obj:
            yield from _walk(v, key, f"{path}[]")
    elif isinstance(obj, str) and key not in ID_KEYS:
        yield path, obj


@pytest.mark.data
def test_site_json_has_no_dashes():
    files = sorted(WEB_DATA.rglob("*.json"))
    if not files:
        pytest.skip("web/public/data not built")
    bad = []
    for f in files:
        for path, s in _walk(json.loads(f.read_text())):
            if offending(s):
                bad.append(f"{f.relative_to(WEB_DATA)}{path}: {s[:90]!r}")
    assert not bad, f"{len(bad)} strings with dashes, e.g.\n" + "\n".join(bad[:25])


@pytest.mark.data
def test_workbook_text_has_no_dashes():
    if not WORKBOOK.exists():
        pytest.skip("workbook not built")
    from openpyxl import load_workbook
    from openpyxl.worksheet.formula import ArrayFormula

    from pipeline.export_excel import _QUOTED

    wb = load_workbook(WORKBOOK)
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, ArrayFormula):
                    v = v.text
                if not isinstance(v, str):
                    continue
                texts = [m.group(1) for m in _QUOTED.finditer(v)] if v.startswith("=") else [v]
                if any(offending(t) for t in texts):
                    bad.append(f"{ws.title}!{c.coordinate}: {v[:90]!r}")
    assert not bad, f"{len(bad)} workbook cells with dashes, e.g.\n" + "\n".join(bad[:25])
